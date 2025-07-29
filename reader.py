import os
from PIL import Image, UnidentifiedImageError
import pytesseract
from PyPDF2 import PdfReader
from docx import Document
import openpyxl

# Required for image OCR
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

# Extensions
IMAGE_EXTENSIONS = [".jpg", ".jpeg", ".png", ".bmp", ".webp", ".gif"]
CODE_EXTENSIONS = [".py", ".java", ".cpp", ".c", ".js", ".sql"]
SKIP_PREFIXES = ["~$"]

def read_file_content(path):
    """
    Read and extract content based on file extension.
    Return None if file is unsupported, corrupted or skipped.
    """
    try:
        filename = os.path.basename(path)
        ext = os.path.splitext(path)[1].lower()

        # Skip temp/lock files like ~$doc.docx
        if any(filename.startswith(pfx) for pfx in SKIP_PREFIXES):
            return None

        if ext == ".txt":
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

        elif ext == ".pdf":
            reader = PdfReader(path)
            return "\n".join([page.extract_text() or "" for page in reader.pages])

        elif ext == ".docx":
            doc = Document(path)
            return "\n".join([p.text for p in doc.paragraphs])

        elif ext in (".xlsx", ".xls"):
            wb = openpyxl.load_workbook(path, data_only=True)
            content = ""
            for sheet in wb:
                for row in sheet.iter_rows(values_only=True):
                    content += " ".join(str(cell) if cell else "" for cell in row) + "\n"
            return content

        elif ext in CODE_EXTENSIONS:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

        elif ext == ".db":
            return f"[Database File: {os.path.basename(path)}]"

        elif ext in IMAGE_EXTENSIONS:
            try:
                img = Image.open(path)
                text = pytesseract.image_to_string(img)
                return f"[Image: {os.path.basename(path)}]\n{text.strip()}"
            except UnidentifiedImageError:
                return None

    except Exception:
        return None

    return None